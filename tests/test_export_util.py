"""Excel 导出工具测试。"""

import openpyxl

from app.utils import export_util


def test_export_sheets_creates_named_worksheets(tmp_path, monkeypatch):
    """多类原始证据写入同一工作簿的独立工作表。"""
    monkeypatch.setattr(export_util, "EXPORT_DIR", str(tmp_path))

    url = export_util.export_sheets_to_excel(
        {
            "小区汇总": [{"cgi": "a"}],
            "扩展明细": [
                {
                    "cgi": "a",
                    "hour_detail": [{"hour": 22, "low_flow_pct": 90.0}],
                }
            ],
            "空工作表": [],
            "收缩明细": [{"cgi": "a", "prb_increase": 10.0}],
        },
        prefix="batch_analysis",
    )

    workbook_path = next(tmp_path.glob("*.xlsx"))
    workbook = openpyxl.load_workbook(workbook_path)
    assert workbook.sheetnames == ["小区汇总", "扩展明细", "收缩明细"]
    assert '"low_flow_pct": 90.0' in workbook["扩展明细"]["B2"].value
    assert workbook["收缩明细"]["B2"].value == 10.0
    for worksheet in workbook.worksheets:
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref == worksheet.dimensions
        assert worksheet["A1"].font.bold is True
    assert url is not None


def test_export_sheets_streams_rows_without_pandas_dataframe(tmp_path, monkeypatch):
    """批量多表导出不得再构造高开销的 pandas DataFrame。"""
    monkeypatch.setattr(export_util, "EXPORT_DIR", str(tmp_path))

    def fail_dataframe(*args, **kwargs):
        raise AssertionError("multi-sheet export must stream rows directly")

    monkeypatch.setattr(export_util.pd, "DataFrame", fail_dataframe)

    url = export_util.export_sheets_to_excel(
        {
            "小区汇总": [
                {"cgi": "a"},
                {"cgi": "b", "cell_name": "测试小区"},
            ],
        },
        prefix="batch_analysis",
    )

    workbook_path = next(tmp_path.glob("*.xlsx"))
    workbook = openpyxl.load_workbook(workbook_path, read_only=True)
    worksheet = workbook["小区汇总"]
    assert url is not None
    assert worksheet["A1"].value == "CGI"
    assert worksheet["B1"].value == "小区名称"
    assert worksheet["A2"].value == "a"
    assert worksheet["B2"].value is None
    assert worksheet["A3"].value == "b"
    assert worksheet["B3"].value == "测试小区"


def test_v14_fields_have_stable_chinese_column_names():
    mapping = export_util.DEFAULT_COLUMN_MAPPING

    assert mapping["deploy_hours_continuous"] == "含已休眠连续部署时段(含扩展时段)"
    assert mapping["prb_increase"] == "邻区PRB抬升量"
    assert mapping["main_site_type"] == "主小区站型"
    assert mapping["around_site_type"] == "邻区站型"
